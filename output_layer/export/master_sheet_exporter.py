# output_layer/export/master_sheet_exporter.py
"""
MasterSheetExporter
===================
Produces an Excel workbook (.xlsx) matching the Master sheet format used by
TKM College of Engineering exam administration.

Output sheets (in order)
------------------------
1. Classroom  – room reference list
2. Master     – one row per group-allocation-in-room (correct hold/row status)
3. <room_no>  – one invigilator hall sheet per used classroom

Status logic
------------
  "firstrow" / "secondrow" / "lastrow" — group starts (or continues within)
      the same room for the first time we see it.
  "firsthold" / "secondhold" / "lasthold" — group was already seen in an
      earlier room (overflow / continuation).

The key insight: all streams belonging to the same group in the *same* room
are normal "row" labels.  Only when the group appears in a *different* room
after already being seen elsewhere does it get a "hold" label.
"""

import logging
import re
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side,
)
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ── Shared styles ──────────────────────────────────────────────────────────────
_HDR_FILL = PatternFill("solid", fgColor="2C3E50")
_HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
_BOLD     = Font(bold=True)
_CENTER   = Alignment(horizontal="center", vertical="center")
_WRAP_CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN     = Border(
    left   = Side(style="thin"),
    right  = Side(style="thin"),
    top    = Side(style="thin"),
    bottom = Side(style="thin"),
)

_STATUS_FILL = {
    "firstrow":   PatternFill("solid", fgColor="D9EAD3"),   # green
    "secondrow":  PatternFill("solid", fgColor="CFE2F3"),   # blue
    "lastrow":    PatternFill("solid", fgColor="FFF2CC"),   # yellow
    "firsthold":  PatternFill("solid", fgColor="F4CCCC"),   # pink
    "secondhold": PatternFill("solid", fgColor="FCE5CD"),   # peach
    "lasthold":   PatternFill("solid", fgColor="EAD1DC"),   # mauve
}

_STREAM_TO_ROW  = {"A": "firstrow",  "B": "secondrow",  "C": "lastrow"}
_STREAM_TO_HOLD = {"A": "firsthold", "B": "secondhold", "C": "lasthold"}


# ══════════════════════════════════════════════════════════════════════════════
# Roll-number utilities
# ══════════════════════════════════════════════════════════════════════════════

def _roll_parts(roll: str):
    """
    Split 'B23CEA01' → ('B23CEA', 1).
    Returns (None, None) if the format doesn't match.
    """
    m = re.match(r"^([A-Za-z]+\d+[A-Za-z]+)0*(\d+)$", roll.strip())
    if m:
        return m.group(1), int(m.group(2))
    return None, None


def _compact_nums(nums: list[int]) -> str:
    """[1,2,3,5,6] → '1-3,5-6'"""
    if not nums:
        return ""
    nums = sorted(nums)
    segs: list[str] = []
    start = end = nums[0]
    for n in nums[1:]:
        if n == end + 1:
            end = n
        else:
            segs.append(f"{start}-{end}" if start != end else str(start))
            start = end = n
    segs.append(f"{start}-{end}" if start != end else str(start))
    return ",".join(segs)


def build_comma_separated(students) -> str:
    """
    Compact roll-number list into a range string.

    Strategy
    --------
    Group students by their roll-number prefix (e.g. 'B23CEA', 'B23CEB').
    For each prefix, compact consecutive numbers into ranges.
    If there is only one prefix → emit bare numbers: '1-15,17-30'
    If there are multiple prefixes (merged A+B group) → prefix each block:
        'A:1-15 / B:1-8'  where A/B is the last letter of the prefix

    Examples
    --------
    [B23CEA01..B23CEA15]                       → '1-15'
    [B23CEA01-05, B23CEA07]                    → '1-5,7'
    [B22ARA01-40, B22ARB01-34] (merged)        → 'A:1-40 / B:1-34'
    [B22ARA32-40, B22ARB01-08] (slice crossing) → 'A:32-40 / B:1-8'
    """
    if not students:
        return ""

    rolls = [s.roll_no for s in students if s.roll_no]
    if not rolls:
        return ""

    # Group by prefix, preserving original order of first appearance
    from collections import OrderedDict
    prefix_groups: OrderedDict[str, list[int]] = OrderedDict()
    no_prefix: list[str] = []

    for roll in rolls:
        prefix, num = _roll_parts(roll)
        if prefix is None:
            no_prefix.append(roll)
            continue
        if prefix not in prefix_groups:
            prefix_groups[prefix] = []
        prefix_groups[prefix].append(num)

    parts: list[str] = []

    # No-prefix rolls (unusual) come first
    parts.extend(no_prefix)

    if len(prefix_groups) == 1:
        # Single prefix — emit bare numbers only
        nums = list(prefix_groups.values())[0]
        parts.append(_compact_nums(nums))
    else:
        # Multiple prefixes — label each block with the last letter of the prefix
        # e.g. 'B22ARA' → label 'A', 'B22ARB' → label 'B'
        for prefix, nums in prefix_groups.items():
            label = prefix[-1] if prefix else prefix   # last char, e.g. 'A' or 'B'
            parts.append(f"{label}:{_compact_nums(nums)}")

    return " / ".join(p for p in parts if p)



# ══════════════════════════════════════════════════════════════════════════════
# Status (hold / row) logic
# ══════════════════════════════════════════════════════════════════════════════

def _build_master_rows(context) -> list[dict]:
    """
    Iterate room_allocations → for each non-empty stream, produce one Master row.

    Hold detection
    --------------
    Track the first room_no where each group_id is seen.
    • Same room_no on a subsequent stream slot → still "row" (it's the same block).
    • Different room_no for a group already seen  → "hold" (overflow continuation).
    """
    rows: list[dict] = []
    group_first_room: dict[str, str] = {}   # group_id → first room_no

    for room in context.room_allocations:
        if not room.allocations:
            continue

        room_no = room.classroom.room_no

        for stream_key in ["A", "B", "C"]:
            alloc = room.streams.get(stream_key)
            if alloc is None:
                continue

            group    = alloc.group
            students = alloc.students
            gid      = group.group_id

            # Determine status
            if gid not in group_first_room:
                group_first_room[gid] = room_no
                status = _STREAM_TO_ROW[stream_key]
            elif group_first_room[gid] == room_no:
                # Same room — different stream — still a "row" (not a hold)
                status = _STREAM_TO_ROW[stream_key]
            else:
                # Spilled into a new room → hold
                status = _STREAM_TO_HOLD[stream_key]

            branch     = f"{group.department}{group.semester}{group.section}"
            start_roll = students[0].roll_no  if students else ""
            end_roll   = students[-1].roll_no if students else ""

            rows.append({
                "Branch":            branch,
                "Count":             len(students),
                "Classname":         room_no,
                "Status":            status,
                "Start Roll Number": start_roll,
                "End Roll Number":   end_roll,
                "Comma Separated":   build_comma_separated(students),
                "Subject":           group.subject_name,
            })

    return rows


# ══════════════════════════════════════════════════════════════════════════════
# Sheet writers
# ══════════════════════════════════════════════════════════════════════════════

def _auto_width(ws, max_w: int = 50) -> None:
    for col in ws.columns:
        mx = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(mx + 3, max_w)


def _write_classroom_sheet(wb, classrooms) -> None:
    ws = wb.create_sheet("Classroom")
    headers = ["Sl No", "Hall No", "No Of Benches", "No of student"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font      = _HDR_FONT
        cell.fill      = _HDR_FILL
        cell.alignment = _CENTER

    for sl, c in enumerate(classrooms, 1):
        benches = c.rows * c.benches_per_row
        ws.append([sl, c.room_no, benches, benches * c.seats_per_bench])

    _auto_width(ws, max_w=20)
    ws.freeze_panes = "A2"


def _write_master_sheet(wb, master_rows: list[dict]) -> None:
    ws = wb.create_sheet("Master")
    COLS = [
        "Branch", "Count", "Classname", "Status",
        "Start Roll Number", "End Roll Number",
        "Comma Separated", "Subject",
    ]
    ws.append(COLS)
    for cell in ws[1]:
        cell.font      = _HDR_FONT
        cell.fill      = _HDR_FILL
        cell.alignment = _CENTER
        cell.border    = _THIN

    for row in master_rows:
        ws.append([row[c] for c in COLS])
        fill = _STATUS_FILL.get(row["Status"], PatternFill("solid", fgColor="FFFFFF"))
        for cell in ws[ws.max_row]:
            cell.fill      = fill
            cell.border    = _THIN
            cell.alignment = _CENTER

    # Fixed column widths
    for col_letter, width in zip("ABCDEFGH", [12, 7, 12, 14, 18, 18, 30, 50]):
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 20
    logger.info("Sheet 'Master' written (%d rows)", len(master_rows))


def _write_hall_sheets(wb, context, exam_info: dict) -> None:
    used_rooms = [r for r in context.room_allocations if r.allocations]

    for sl_no, room in enumerate(used_rooms, start=1):
        room_no    = room.classroom.room_no
        sheet_name = room_no[:31]
        if sheet_name in wb.sheetnames:
            sheet_name = f"{sheet_name[:28]}_{sl_no}"

        ws = wb.create_sheet(sheet_name)

        college   = exam_info.get("college_name", "")
        exam_name = exam_info.get("exam_name", "")
        exam_date = exam_info.get("exam_date", "")
        session   = exam_info.get("session", "")

        # Row 1: college name (col A, spans A:E) + serial number (col F)
        ws.append([
            f"                    {college}",
            "", "", "", "",
            f"Sl No: {sl_no}",
        ])
        ws.merge_cells("A1:E1")
        ws["A1"].font      = Font(bold=True, size=11)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["F1"].alignment = Alignment(horizontal="right")

        # Row 2: institution type
        ws.append(["Government Aided and Autonomous"])
        ws.merge_cells("A2:F2")
        ws["A2"].alignment = _CENTER

        # Row 3: exam name
        ws.append([exam_name])
        ws.merge_cells("A3:F3")
        ws["A3"].alignment = _CENTER

        # Row 4: date + session + hall
        ws.append([
            f"Date: {exam_date}", "",
            session, "",
            f"Hall No.: {room_no}",
        ])

        # Row 5: column headers
        ws.append(["Branch", "Subject", "Roll nos.", "No.of Stds", "Absentees nos.", "Nos."])
        for cell in ws[5]:
            cell.font      = _BOLD
            cell.fill      = PatternFill("solid", fgColor="D9D9D9")
            cell.border    = _THIN
            cell.alignment = _CENTER

        # Data rows
        total_students = 0
        for stream_key in ["A", "B", "C"]:
            alloc = room.streams.get(stream_key)
            if alloc is None:
                continue

            group    = alloc.group
            students = alloc.students
            count    = len(students)
            total_students += count

            branch   = f"{group.department}{group.semester}{group.section}"
            roll_str = build_comma_separated(students)

            ws.append([branch, group.subject_name, roll_str, count, "", ""])
            for cell in ws[ws.max_row]:
                cell.border    = _THIN
                cell.alignment = _WRAP_CTR

        # Signature row
        sig_row = ws.max_row + 1
        ws.append(["Name & Signature of Invigilator", "", "", "", "Verifier", ""])
        ws.merge_cells(f"A{sig_row}:D{sig_row}")
        ws[f"A{sig_row}"].font = _BOLD

        # Total students
        ws.append([f"Total Students: {total_students}"])
        ws[f"A{ws.max_row}"].font = _BOLD

        # Column widths
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 10


# ══════════════════════════════════════════════════════════════════════════════
# Public exporter class
# ══════════════════════════════════════════════════════════════════════════════

class MasterSheetExporter:
    """
    Exports the allocation context to an Excel workbook in Master sheet format.

    Usage
    -----
    MasterSheetExporter().export(
        context    = context,
        classrooms = classrooms,
        exam_info  = {
            "college_name": "TKM College of Engineering, Kollam-5",
            "exam_name":    "Second Series Exam, March 2026",
            "exam_date":    "12-03-2026",
            "session":      "FN",
        },
        filepath   = "output/master_12-03-2026_FN.xlsx",
    )
    """

    def export(self, context, classrooms, exam_info: dict,
               filepath: str | Path) -> None:
        filepath = Path(filepath)
        filepath.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        _write_classroom_sheet(wb, classrooms)
        master_rows = _build_master_rows(context)
        _write_master_sheet(wb, master_rows)
        _write_hall_sheets(wb, context, exam_info)

        wb.save(filepath)
        used = sum(1 for r in context.room_allocations if r.allocations)
        logger.info(
            "Master Sheet Excel saved → %s  (%d master rows, %d hall sheets)",
            filepath, len(master_rows), used,
        )
