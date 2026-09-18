from pathlib import Path
from openpyxl import load_workbook
from engine.models.student import Student


class StudentLoader:
    def __init__(self, student_file: str, timetable_file: str):
        self.student_file = Path(student_file)
        self.timetable_file = Path(timetable_file)
        self._unknown_counter = 0  # For generating placeholder IDs
        self._warnings = []  # Collect all warnings to report at end

    def load(self) -> list[Student]:
        timetable = self._load_timetable()
        workbook = load_workbook(self.student_file, data_only=True)
        students: list[Student] = []
        self._unknown_counter = 0
        self._warnings = []

        for sheet_name in workbook.sheetnames:
            if sheet_name.lower() == "master overview":
                continue

            worksheet = workbook[sheet_name]
            department = self._get_department(sheet_name)
            semester = self._get_semester(sheet_name)
            section = self._get_section(sheet_name)

            exams = timetable.get((department, semester), [])
            if not exams:
                continue

            for row_num, row in enumerate(worksheet.iter_rows(min_row=7, values_only=True), start=7):
                if row[0] is None:
                    break

                # Safely extract row data with bounds checking
                row_data = self._safe_extract_row(row, row_num, section, department)

                # Log warnings for auto-fixes
                if row_data["warnings"]:
                    self._warnings.extend(row_data["warnings"])

                for exam in exams:
                    students.append(
                        Student(
                            register_no=row_data["register_no"],
                            name=row_data["name"],
                            department=department,
                            semester=semester,
                            section=section,
                            subject_code=exam["subject_code"],
                            subject_name=exam["subject_name"],
                            exam_date=str(exam["exam_date"] or "").strip(),
                            session=str(exam["session"] or "").strip(),
                            roll_no=row_data["roll_no"],
                        )
                    )

        workbook.close()

        # Report all warnings at the end
        self._report_warnings()

        return students

    def _safe_extract_row(self, row, row_num, section, department):
        """Safely extract student data from row, auto-fixing issues with alerts."""
        warnings = []
        result = {
            "roll_no": "",
            "register_no": "",
            "name": "",
            "warnings": warnings,
        }

        # Extract roll_no (column index 2)
        try:
            if len(row) > 2 and row[2] is not None:
                result["roll_no"] = str(row[2]).strip()
            else:
                warnings.append(
                    f"  ⚠ Row {row_num}: Missing roll_no, auto-fixed"
                )
        except Exception:
            result["roll_no"] = ""
            warnings.append(f"  ⚠ Row {row_num}: Invalid roll_no format, auto-fixed")

        # Extract register_no (column index 3)
        try:
            if len(row) > 3 and row[3] is not None:
                base_reg_no = str(row[3]).strip()
                if not base_reg_no or base_reg_no.lower() in ("nan", "none", ""):
                    raise ValueError("Empty register number")
            else:
                raise ValueError("Missing register number column")
        except (ValueError, IndexError):
            self._unknown_counter += 1
            base_reg_no = f"UNKNOWN_{self._unknown_counter:04d}"
            warnings.append(
                f"  ⚠ Row {row_num}: Missing register_no '{row[3] if len(row) > 3 else 'N/A'}', "
                f"generated '{base_reg_no}'"
            )

        # Add section disambiguation
        result["register_no"] = f"{base_reg_no}-{section}" if section else base_reg_no

        # Extract name (column index 4)
        try:
            if len(row) > 4 and row[4] is not None:
                name = str(row[4]).strip()
                if not name or name.lower() in ("nan", "none", ""):
                    raise ValueError("Empty name")
                result["name"] = name
            else:
                raise ValueError("Missing name column")
        except (ValueError, IndexError):
            result["name"] = "UNKNOWN_STUDENT"
            warnings.append(f"  ⚠ Row {row_num}: Missing name, filled with 'UNKNOWN_STUDENT'")

        return result

    def _report_warnings(self):
        """Report all collected warnings at the end of loading."""
        if not self._warnings:
            return

        print("\n" + "=" * 65)
        print(" ⚠ DATA AUTO-FIX ALERTS (Student Loader)")
        print("=" * 65)
        print(f" Total issues found: {len(self._warnings)}")
        print("-" * 65)

        # Group warnings by type
        reg_warnings = [w for w in self._warnings if "register_no" in w]
        name_warnings = [w for w in self._warnings if "name" in w and "register" not in w]
        roll_warnings = [w for w in self._warnings if "roll_no" in w]

        if reg_warnings:
            print(f"\n 📋 Register Number Issues ({len(reg_warnings)}):")
            for w in reg_warnings[:10]:  # Show first 10
                print(w)
            if len(reg_warnings) > 10:
                print(f"  ... and {len(reg_warnings) - 10} more")

        if name_warnings:
            print(f"\n 📋 Name Issues ({len(name_warnings)}):")
            for w in name_warnings[:10]:
                print(w)
            if len(name_warnings) > 10:
                print(f"  ... and {len(name_warnings) - 10} more")

        if roll_warnings:
            print(f"\n 📋 Roll Number Issues ({len(roll_warnings)}):")
            for w in roll_warnings[:10]:
                print(w)
            if len(roll_warnings) > 10:
                print(f"  ... and {len(roll_warnings) - 10} more")

        print("\n" + "=" * 65)
        print(" ℹ Students with generated IDs will be skipped during allocation")
        print("   unless manual correction is made in source data.")
        print("=" * 65 + "\n")

    def _load_timetable(self) -> dict[tuple[str, int], list[dict]]:
        workbook = load_workbook(self.timetable_file, data_only=True)
        worksheet = workbook.active
        timetable: dict[tuple[str, int], list[dict]] = {}
        all_btech_depts = ["CE", "ME", "EE", "EC", "CS", "CH", "EL"]

        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None or row[5] is None:
                continue

            program = str(row[0]).strip()
            sem_str = str(row[1]).strip().upper().replace("S", "")
            if not sem_str.isdigit():
                continue
            semester = int(sem_str)

            branch_raw = str(row[5]).strip().upper()
            if program == "B Arch":
                target_depts = ["B.ARCH"]
            elif branch_raw == "ALL BRANCHES":
                target_depts = all_btech_depts
            elif "/" in branch_raw:
                target_depts = [b.strip() for b in branch_raw.split("/")]
            else:
                target_depts = [branch_raw]

            exam_data = {
                "exam_date": str(row[2]).strip() if row[2] else "",
                "session": str(row[3]).strip() if row[3] else "",
                "subject_name": str(row[6]).strip() if row[6] else "",
                "subject_code": str(row[7]).strip() if row[7] else "",
            }

            for dept in target_depts:
                key = (dept, semester)
                if key not in timetable:
                    timetable[key] = []
                timetable[key].append(exam_data)

        workbook.close()
        return timetable

    @staticmethod
    def _get_department(sheet_name: str) -> str:
        raw = sheet_name.split()[0].upper().replace(".", "")
        if raw == "BARCH":
            return "B.ARCH"
        if raw == "EEE":
            return "EE"
        return raw

    @staticmethod
    def _get_semester(sheet_name: str) -> int:
        start = sheet_name.find("(S")
        end = sheet_name.find(")", start)
        return int(sheet_name[start + 2:end])

    @staticmethod
    def _get_section(sheet_name: str) -> str:
        parts = sheet_name.split()
        if len(parts) >= 3 and not parts[-2].startswith("(S"):
            return parts[-2]
        return ""