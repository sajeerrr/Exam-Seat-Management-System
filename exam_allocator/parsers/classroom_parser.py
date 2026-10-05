"""
Excel classroom parser.
"""

from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from openpyxl import load_workbook
import pandas as pd

class ClassroomExcelParseError(Exception):
    """Raised when the classroom Excel workbook cannot be parsed."""

@dataclass
class ClassroomRecord:
    room_number: str
    capacity: int
    benches: int
    building: str

@dataclass
class ClassroomExtractionResult:
    source_file: str
    classrooms: list[ClassroomRecord]
    issues: list[str]

DEFAULT_CAPACITY = 45
DEFAULT_BENCHES = 15
DEFAULT_BUILDING = "Main Building"

def parse_classroom_excel(
    excel_path: str,
    default_capacity: int = DEFAULT_CAPACITY,
    default_benches: int = DEFAULT_BENCHES,
    default_building: str = DEFAULT_BUILDING,
) -> ClassroomExtractionResult:
    path = Path(excel_path)

    if not path.exists():
        raise ClassroomExcelParseError(f"File not found: {excel_path}")

    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise ClassroomExcelParseError(
            f"Unsupported file type: {path.suffix}"
        )

    classrooms: list[ClassroomRecord] = []
    issues: list[str] = []
    seen = set()

    # Strategy 1: openpyxl multi-column pair scanning (e.g. multi-block sheets)
    try:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        
        hall_col_pairs = []
        for r_idx in range(min(10, len(rows))):
            row = rows[r_idx]
            for c_idx in range(len(row) - 1):
                c_val = str(row[c_idx] or "").strip().lower()
                c_next = str(row[c_idx + 1] or "").strip().lower()
                if "sl" in c_val and any(k in c_next for k in ["hall", "room", "class"]):
                    hall_col_pairs.append((c_idx, c_idx + 1))
        
        if hall_col_pairs:
            current_building = {}
            for row in rows:
                for sl_col, hall_col in hall_col_pairs:
                    val = row[sl_col]
                    if val and isinstance(val, str) and any(k in val for k in ["Block", "Studio", "Building"]):
                        current_building[hall_col] = val.strip()
                    elif val and isinstance(val, (int, float)) and int(val) > 0:
                        room_val = row[hall_col]
                        if room_val:
                            r_str = str(room_val).strip()
                            if r_str and r_str.lower() not in {"nan", "none"} and not r_str.lower().startswith("sl"):
                                if r_str not in seen:
                                    seen.add(r_str)
                                    b_str = current_building.get(hall_col, default_building)
                                    classrooms.append(
                                        ClassroomRecord(
                                            room_number=r_str,
                                            capacity=default_capacity,
                                            benches=default_benches,
                                            building=b_str
                                        )
                                    )
    except Exception:
        pass

    # Strategy 2: If no classrooms found via strategy 1, use pandas table parsing
    if not classrooms:
        try:
            df = pd.read_excel(path)
            df.columns = df.columns.astype(str).str.lower().str.strip()
            
            hall_cols = [c for c in df.columns if any(x in c for x in ["hall", "room", "class", "name"])]
            cap_cols = [c for c in df.columns if any(x in c for x in ["student", "capacit"])]
            bench_cols = [c for c in df.columns if any(x in c for x in ["bench"])]
            
            if not hall_cols:
                if len(df.columns) > 1:
                    hall_cols = [df.columns[1]]
                elif len(df.columns) == 1:
                    hall_cols = [df.columns[0]]
                else:
                    raise ClassroomExcelParseError("Could not find a column for 'Hall No' or 'Room No'.")
            
            hall_col = hall_cols[0]
            
            for idx, row in df.iterrows():
                room_val = row[hall_col]
                if pd.isna(room_val):
                    continue
                    
                room_str = str(room_val).strip()
                if not room_str or room_str.lower() in {"nan", "none"}:
                    continue
                    
                if room_str.lower() in {"class room", "classroom", "room", "room number", "hall no", "cctv availabe class room list", "main block", "sl no"}:
                    continue
                    
                if room_str in seen:
                    issues.append(f"Duplicate classroom '{room_str}'. Skipped.")
                    continue
                seen.add(room_str)
                    
                capacity = default_capacity
                if cap_cols:
                    val = row[cap_cols[0]]
                    if pd.notna(val):
                        try:
                            capacity = int(float(val))
                        except ValueError:
                            pass

                benches = default_benches
                if bench_cols:
                    val = row[bench_cols[0]]
                    if pd.notna(val):
                        try:
                            benches = int(float(val))
                        except ValueError:
                            pass

                classrooms.append(
                    ClassroomRecord(
                        room_number=room_str,
                        capacity=capacity,
                        benches=benches,
                        building=default_building
                    )
                )
        except Exception as exc:
            raise ClassroomExcelParseError(f"Could not parse classroom file: {exc}") from exc

    if not classrooms:
        raise ClassroomExcelParseError("No valid classrooms were found in the workbook.")

    return ClassroomExtractionResult(
        source_file=path.name,
        classrooms=classrooms,
        issues=issues,
    )
